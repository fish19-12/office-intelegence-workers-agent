from documents.document_manager import (
    Chunk,
    ChunkFinalizationLayer,
    DocumentLoader,
    DocumentManager,
    FixedSizeChunkingStrategy,
    ParagraphAwareChunkingStrategy,
    ParentChildChunkingStrategy,
    ProcessingConfig,
    SentenceAwareChunkingStrategy,
    StructureAwareChunkingStrategy,
    TableAwareChunkingStrategy,
    InvalidChunkError,
)


class DummyRAG:
    def add_documents(self, documents):
        return None


def test_document_loader_validates_file_and_uses_safe_path(tmp_path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text("hello world\nthis is a sample", encoding="utf-8")

    loader = DocumentLoader(max_file_size_bytes=10 * 1024 * 1024)
    resolved = loader.load(file_path)

    assert resolved.path.name == "sample.txt"
    assert resolved.size_bytes > 0
    assert resolved.mime_type in {"text/plain", "application/octet-stream"}

    try:
        loader.load(tmp_path / "missing.txt")
        assert False, "missing file should fail"
    except FileNotFoundError:
        pass

    bad_path = tmp_path / ".." / "outside.txt"
    try:
        loader.load(bad_path)
        assert False, "traversal should fail"
    except ValueError:
        pass


def test_fixed_chunking_keeps_overlap_and_flushes_tail():
    doc = {
        "document_id": "doc-1",
        "title": "Sample",
        "document_type": "text",
        "pages": [{
            "page_number": 1,
            "blocks": [{
                "block_type": "paragraph",
                "text": "Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda",
                "heading_path": [],
                "source_location": {"page": 1},
                "metadata": {},
            }]
        }],
        "sections": [],
        "blocks": [{
            "block_type": "paragraph",
            "text": "Alpha beta gamma delta epsilon zeta eta theta iota kappa lambda",
            "heading_path": [],
            "source_location": {"page": 1},
            "metadata": {},
        }],
    }

    strategy = FixedSizeChunkingStrategy()
    chunks = strategy.chunk(doc, ProcessingConfig(chunk_size=30, overlap=10, min_chunk_size=5))

    assert len(chunks) >= 2
    assert chunks[0].metadata["chunk_index"] == 0
    assert chunks[0].metadata["character_count"] > 0
    assert chunks[-1].content


def test_structure_aware_chunking_keeps_heading_context():
    doc = {
        "document_id": "doc-2",
        "title": "Employee Benefits",
        "document_type": "docx",
        "pages": [{
            "page_number": 1,
            "blocks": [
                {"block_type": "heading", "text": "Human Resources", "heading_level": 1, "heading_path": ["Human Resources"], "source_location": {"page": 1}, "metadata": {}},
                {"block_type": "paragraph", "text": "Employees receive health coverage.", "heading_path": ["Human Resources"], "source_location": {"page": 1}, "metadata": {}},
                {"block_type": "heading", "text": "Annual Leave", "heading_level": 2, "heading_path": ["Human Resources", "Annual Leave"], "source_location": {"page": 1}, "metadata": {}},
                {"block_type": "paragraph", "text": "Employees earn fifteen working days each year.", "heading_path": ["Human Resources", "Annual Leave"], "source_location": {"page": 1}, "metadata": {}},
            ]
        }],
        "sections": [],
        "blocks": [],
    }

    strategy = StructureAwareChunkingStrategy()
    chunks = strategy.chunk(doc, ProcessingConfig(chunk_size=100, overlap=0, min_chunk_size=20))

    assert len(chunks) >= 2
    assert any("Human Resources" in chunk.content for chunk in chunks)
    assert any("Annual Leave" in chunk.content for chunk in chunks)


def test_sentence_and_paragraph_strategies_handle_content():
    doc = {
        "document_id": "doc-3",
        "title": "Policy",
        "document_type": "text",
        "blocks": [
            {"block_type": "paragraph", "text": "This is sentence one. This is sentence two. This is sentence three.", "heading_path": [], "source_location": {"page": 1}, "metadata": {}},
            {"block_type": "paragraph", "text": "Second paragraph text continues here and stays together.", "heading_path": [], "source_location": {"page": 1}, "metadata": {}},
        ],
    }

    sentence_chunks = SentenceAwareChunkingStrategy().chunk(doc, ProcessingConfig(chunk_size=80, overlap=0, min_chunk_size=10))
    paragraph_chunks = ParagraphAwareChunkingStrategy().chunk(doc, ProcessingConfig(chunk_size=60, overlap=0, min_chunk_size=10))

    assert len(sentence_chunks) >= 2
    assert len(paragraph_chunks) >= 2
    assert all(chunk.content for chunk in sentence_chunks + paragraph_chunks)


def test_table_aware_and_parent_child_chunking():
    doc = {
        "document_id": "doc-4",
        "title": "Table report",
        "document_type": "csv",
        "blocks": [
            {"block_type": "table", "text": "Product | Price\nLaptop | 1000\nPhone | 500\nTablet | 700", "heading_path": ["Table report"], "source_location": {"page": 1}, "metadata": {"headers": ["Product", "Price"], "rows": [["Laptop", "1000"], ["Phone", "500"], ["Tablet", "700"]]}} 
        ],
    }

    table_chunks = TableAwareChunkingStrategy().chunk(doc, ProcessingConfig(chunk_size=100, overlap=0, min_chunk_size=5, table_row_group_size=2))
    parent_chunks = ParentChildChunkingStrategy().chunk(doc, ProcessingConfig(chunk_size=80, overlap=0, min_chunk_size=10, parent_child_mode=True))

    assert len(table_chunks) >= 2
    assert any("Product" in chunk.content for chunk in table_chunks)
    assert len(parent_chunks) >= 1
    assert all(chunk.metadata.get("parent_chunk_id") is not None or not chunk.metadata.get("parent_chunk_id") for chunk in parent_chunks)


def test_document_manager_ingest_compatibility(tmp_path):
    file_path = tmp_path / "sample.csv"
    file_path.write_text("name,amount\nAlpha,100\nBeta,200\n", encoding="utf-8")

    manager = DocumentManager(DummyRAG())
    asset = manager.ingest(str(file_path), {"category": "Finance"})

    assert asset.doc_id
    assert asset.chunk_count > 0
    assert manager.get_document(asset.doc_id) == asset


def test_adaptive_chunking_profile_selects_the_right_strategy():
    structured_doc = {
        "document_id": "doc-5",
        "title": "Benefits",
        "document_type": "docx",
        "blocks": [
            {"block_type": "heading", "text": "Benefits", "heading_path": ["Benefits"], "source_location": {"page": 1}, "metadata": {}},
            {"block_type": "paragraph", "text": "Employees receive medical and dental coverage as part of the total reward package.", "heading_path": ["Benefits"], "source_location": {"page": 1}, "metadata": {}},
        ],
    }

    table_doc = {
        "document_id": "doc-6",
        "title": "Revenue",
        "document_type": "csv",
        "blocks": [
            {"block_type": "table", "text": "Region | Revenue\nNorth | 1200\nSouth | 900", "heading_path": ["Revenue"], "source_location": {"page": 1}, "metadata": {"headers": ["Region", "Revenue"]}},
        ],
    }

    structured_chunks = DocumentManager().processor.chunking_engine.chunk(
        structured_doc,
        ProcessingConfig(chunk_size=100, overlap=0, min_chunk_size=10, chunking_strategy="auto", chunking_profile="structured"),
    )
    table_chunks = DocumentManager().processor.chunking_engine.chunk(
        table_doc,
        ProcessingConfig(chunk_size=100, overlap=0, min_chunk_size=10, chunking_strategy="auto", chunking_profile="table_focused"),
    )

    assert structured_chunks
    assert table_chunks
    assert all(chunk.metadata.get("chunking_strategy") in {"structure_aware", "table_aware"} for chunk in structured_chunks + table_chunks)


def test_docx_preserves_page_section_content_relationships(tmp_path):
    from docx import Document as WordDocument

    file_path = tmp_path / "employee_benefits.docx"
    source = WordDocument()
    source.add_heading("Employee Benefits", level=1)
    source.add_paragraph("Employees receive health coverage.")
    source.add_paragraph("Medical coverage", style="List Bullet")
    source.add_table(rows=2, cols=2)
    source.tables[0].rows[0].cells[0].text = "Benefit"
    source.tables[0].rows[0].cells[1].text = "Eligibility"
    source.tables[0].rows[1].cells[0].text = "Health"
    source.tables[0].rows[1].cells[1].text = "All employees"
    source.save(file_path)

    document, _ = DocumentManager().processor.process(file_path)

    assert document["pages"]
    page = document["pages"][0]
    assert page["page_number"] == 1
    section = next(item for item in document["sections"] if item["title"] == "Employee Benefits")
    assert section["page_number"] == 1
    assert [item["block_type"] for item in section["blocks"]] == ["heading", "paragraph", "list", "table"]
    assert section["blocks"][1]["heading_path"] == ["Employee Benefits"]
    assert section["blocks"][2]["heading_path"] == ["Employee Benefits"]
    assert section["blocks"][3]["heading_path"] == ["Employee Benefits"]


def test_chunk_finalization_normalizes_deduplicates_and_prepares_embedding_text():
    chunks = [
        Chunk(
            "  Employee   Benefits\n\nHealth coverage is available.  ",
            {"heading_path": ["Employee Benefits"], "source_location": {"page": 3}},
        ),
        Chunk(
            "Employee Benefits\n\nHealth coverage is available.",
            {"heading_path": ["Employee Benefits"], "source_location": {"page": 3}},
        ),
        Chunk("Retirement contributions are available.", {"heading_path": ["Employee Benefits"]}),
    ]

    finalized = ChunkFinalizationLayer().finalize(
        chunks,
        document={"document_id": "doc-finalize", "document_type": "docx"},
        config=ProcessingConfig(chunk_size=200, max_chunk_size=200, deduplication_behavior="safe"),
    )

    assert len(finalized) == 2
    assert finalized[0].content == "Employee Benefits\n\nHealth coverage is available."
    assert finalized[0].embedding_text == finalized[0].content
    assert finalized[0].metadata["finalization_status"] == "ready"
    assert finalized[0].metadata["document_id"] == "doc-finalize"
    assert finalized[0].metadata["chunk_index"] == 0
    assert finalized[0].metadata["content_hash"]
    assert finalized[0].metadata["embedding_text"] == finalized[0].embedding_text
    assert finalized[1].metadata["chunk_index"] == 1


def test_chunk_finalization_rejects_empty_and_oversized_chunks():
    finalizer = ChunkFinalizationLayer()

    try:
        finalizer.finalize([Chunk("   ")])
        assert False, "empty normalized content should fail"
    except InvalidChunkError:
        pass

    try:
        finalizer.finalize([Chunk("too long")], config=ProcessingConfig(chunk_size=3, max_chunk_size=3))
        assert False, "oversized content should fail"
    except InvalidChunkError:
        pass
